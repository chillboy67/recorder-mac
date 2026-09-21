"""
Export writers — the artefacts a user actually reads.

Previously the suite only asserted that the files existed and contained the
transcript. These tests pin the formats themselves (timestamps, speaker labels,
SRT block structure, the JSON contract), because those are what a downstream
tool or a human reads, and a silent format drift would go unnoticed.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from modules import ASRResult, ASRSegment, ASRWord  # noqa: E402
from core import export as E  # noqa: E402


def seg(i, text, start, end, lang="en", conf=-0.25, words=()):
    return ASRSegment(id=i, start=start, end=end, text=text, language=lang,
                      confidence=conf, words=list(words))


def sample_asr() -> ASRResult:
    return ASRResult(
        segments=[
            seg(0, "Hello there", 0.0, 1.5, words=[
                ASRWord("Hello", 0.0, 0.7, -0.1), ASRWord(" there", 0.7, 1.5, -0.4)]),
            seg(1, "I am fine", 3661.25, 3663.0),
        ],
        full_text="Hello there I am fine", language="en", model_used="large-v3",
        audio_duration_sec=3663.0, warnings=["chunk merged"],
    )


# ── timestamp helpers ───────────────────────────────────────────────

@pytest.mark.parametrize("seconds,expected", [
    (0, "0:00"), (5, "0:05"), (65, "1:05"), (3600, "1:00:00"), (3661, "1:01:01"),
])
def test_ts_short(seconds, expected):
    assert E._ts_short(seconds) == expected


@pytest.mark.parametrize("seconds,expected", [
    (0, "00:00:00,000"), (1.5, "00:00:01,500"), (3661.456, "01:01:01,456"),
])
def test_ts_srt_keeps_milliseconds(seconds, expected):
    assert E._ts_srt(seconds) == expected


# ── speaker labels ──────────────────────────────────────────────────

def test_label_is_empty_without_diarization():
    assert E._label(None, 0) == ""
    assert E._label({}, 0) == ""
    assert E._label({1: "candidate"}, 0) == ""      # this segment has no label


def test_label_translates_the_internal_roles():
    labels = {0: "candidate", 1: "examiner", 2: "other", 3: "main"}
    assert [E._label(labels, i) for i in range(4)] == ["考生", "教官", "其他", "主讲"]


def test_label_passes_an_unknown_role_through():
    assert E._label({0: "guest"}, 0) == "guest"


# ── txt ─────────────────────────────────────────────────────────────

def test_txt_is_timestamped_and_labels_the_speaker(tmp_path):
    out = E._txt(sample_asr(), tmp_path, "s", {1: "candidate"})
    assert out.read_text(encoding="utf-8") == (
        "[0:00] Hello there\n[1:01:01] 考生: I am fine")


def test_txt_falls_back_to_full_text_when_there_are_no_segments(tmp_path):
    asr = ASRResult(segments=[], full_text="only text", language="en",
                    model_used="m")
    assert E._txt(asr, tmp_path, "s", None).read_text("utf-8") == "only text"


# ── md ──────────────────────────────────────────────────────────────

def test_md_has_a_metadata_header_and_bold_speakers(tmp_path):
    text = E._md(sample_asr(), tmp_path, "take", {0: "examiner"}).read_text("utf-8")
    assert text.startswith("# take\n")
    assert "- 语言：en" in text and "- 引擎：large-v3" in text
    assert "- 时长：3663s" in text
    assert "## 转写" in text
    assert "**教官** `[0:00]` Hello there" in text
    assert "`[1:01:01]` I am fine" in text          # unlabelled → no bold prefix


# ── srt ─────────────────────────────────────────────────────────────

def test_srt_blocks_are_numbered_and_well_formed(tmp_path):
    text = E._srt(sample_asr(), tmp_path, "s", {1: "candidate"}).read_text("utf-8")
    blocks = text.strip().split("\n\n")
    assert len(blocks) == 2
    assert blocks[0].splitlines() == ["1", "00:00:00,000 --> 00:00:01,500",
                                      "Hello there"]
    assert blocks[1].splitlines() == ["2", "01:01:01,250 --> 01:01:03,000",
                                      "考生: I am fine"]


# ── json ────────────────────────────────────────────────────────────

def test_json_exposes_the_full_segment_contract(tmp_path):
    asr = sample_asr()
    asr.annotations.append({"type": "repeat_arbitration", "verdict": "asr_loop",
                            "original": "no no no no", "folded": False,
                            "segment_id": 0, "start": 0.0, "end": 1.0,
                            "evidence": {"level": 1}})
    data = json.loads(E._json(asr, tmp_path, "s", {1: "candidate"}).read_text("utf-8"))

    assert data["language"] == "en" and data["model"] == "large-v3"
    assert data["duration_sec"] == 3663.0
    assert data["full_text"] == "Hello there I am fine"
    assert data["warnings"] == ["chunk merged"]
    first = data["segments"][0]
    assert first["id"] == 0 and first["start"] == 0.0
    assert first["speaker"] is None                  # no label for segment 0
    assert data["segments"][1]["speaker"] == "candidate"
    assert first["avg_logprob"] == -0.25             # rounded to 3 decimals
    assert first["words"][0] == {"word": "Hello", "start": 0.0, "end": 0.7,
                                "confidence": -0.1}
    # annotations travel with the json export (this is where folded lives)
    assert data["annotations"][0]["folded"] is False
    assert data["annotations"][0]["original"] == "no no no no"


# ── docx / doc ──────────────────────────────────────────────────────

def test_md_into_docx_maps_headings_quotes_and_bullets():
    from docx import Document
    doc = Document()
    E._md_into_docx(doc, "# Title\n\n## Section\n\n### Sub\n\n> quoted\n\n- item\n\nplain\n")
    by_text = {p.text: p.style.name for p in doc.paragraphs if p.text}
    assert by_text["Title"] == "Heading 1"
    assert by_text["Section"] == "Heading 2"
    assert by_text["Sub"] == "Heading 3"
    assert by_text["item"] == "List Bullet"
    assert by_text["plain"] == "Normal"
    assert "quoted" in by_text and by_text["quoted"] != "Normal"


def test_docx_with_a_report_contains_the_report_not_a_second_transcript():
    from docx import Document
    asr = sample_asr()
    doc = E._build_docx_document(asr, "take", None, "# 报告\n\n正文")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "报告" in text and "正文" in text
    assert "Hello there" not in text         # the report already embeds it


def test_docx_without_a_report_is_the_transcript(tmp_path):
    from docx import Document
    out = E._docx(sample_asr(), tmp_path, "take", {0: "examiner"}, None)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "take" in text and "语言 en" in text
    assert "[0:00]" in text and "教官: " in text and "Hello there" in text


def test_doc_raises_a_clear_error_without_textutil(tmp_path, monkeypatch):
    monkeypatch.setattr("shutil.which", lambda name: None)
    with pytest.raises(RuntimeError):
        E._doc(sample_asr(), tmp_path, "s", None, None)


def test_doc_raises_when_textutil_fails(tmp_path, monkeypatch):
    from types import SimpleNamespace
    monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/textutil")
    monkeypatch.setattr("subprocess.run",
                        lambda *a, **kw: SimpleNamespace(returncode=1,
                                                         stderr="boom"))
    with pytest.raises(RuntimeError):
        E._doc(sample_asr(), tmp_path, "s", None, None)


# ── export_all orchestration ────────────────────────────────────────

def test_export_all_ignores_unknown_formats(tmp_path):
    out = E.export_all(sample_asr(), tmp_path, "s", ["txt", "parquet"])
    assert set(out) == {"txt"}


def test_one_failing_writer_does_not_sink_the_others(tmp_path, monkeypatch):
    """`.doc` needs macOS textutil; its absence must not cost the txt export."""
    monkeypatch.setattr("subprocess.run",
                        lambda *a, **kw: __import__("types").SimpleNamespace(
                            returncode=1, stderr="nope"))
    out = E.export_all(sample_asr(), tmp_path, "s", ["txt", "doc", "json"])
    assert "txt" in out and "json" in out
    assert "doc" not in out
    assert out["txt"].exists() and out["json"].exists()


def test_export_all_falls_back_when_the_folder_is_unwritable(tmp_path, monkeypatch):
    """macOS TCC can block the chosen folder; output then goes to data_root()."""
    from core import paths
    monkeypatch.setattr(paths, "data_root", lambda: tmp_path / "fallback")
    blocked = tmp_path / "blocked"
    blocked.write_text("i am a file, not a directory")   # mkdir will fail

    out = E.export_all(sample_asr(), blocked, "s", ["txt"])

    assert out["txt"].parent == tmp_path / "fallback" / "s"
    assert out["txt"].exists()
