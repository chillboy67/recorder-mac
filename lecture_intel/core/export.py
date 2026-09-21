"""
Output writers: txt / md / srt / json, with optional speaker labels.

Self-contained (doesn't depend on the old Exporter's LLM coupling). Timestamps
are always available; speaker labels are emitted only when diarization ran.
"""
from __future__ import annotations

import json
from datetime import timedelta
from pathlib import Path
from typing import Optional

from modules import ASRResult, ASRSegment

from core.i18n import t

_SPEAKER_ZH = {
    "candidate": "考生",
    "examiner": "教官",
    "other": "其他",
    "main": "主讲",
}


# ── Fidelity annotations ─────────────────────────────────────────────────
# The transcript body is never rewritten by this section. It is a report
# *about* the transcript: which spans the three-level arbitration flagged, what
# it decided, and on what evidence. Nothing here is part of the utterance.
FIDELITY_HEADING = "忠实度标注（正文未改写）"
FIDELITY_NOTE = ("正文一律保持说话人的原话。下列片段仅为标注；"
                 "只有课堂模式会把已确认为转写伪影的词级重复折叠，折叠掉的原话记在每一条里。")

_VERDICT_ZH = {
    "asr_loop": "疑似转写伪影",
    "real_speech": "真实重复",
    "uncertain": "判定存疑",
}


def _evidence_brief(ev: dict) -> str:
    """One short, checkable justification drawn from the arbitration evidence."""
    if not ev:
        return ""
    level = ev.get("level")
    if level == 1:
        return (f"L1 词时间轴 {ev.get('span_sec')}s ≪ 应有的 {ev.get('expected_sec')}s"
                "（解码时钟冻结）")
    if level == 2:
        return (f"L2 浊音段 {ev.get('voiced_bursts')} 个 ≥ 0.8×{ev.get('k')}"
                "（每份拷贝都独立发声）")
    if level == 3:
        where = "降噪前原始音频" if ev.get("oracle") == "original" else "管道音频"
        seeds = ev.get("seeds")
        if seeds is None:
            return f"L3 在{where}上重解码失败，证据不足"
        return (f"L3 在{where}上 3 个温度种子复现 {seeds} 份拷贝"
                f"（原文共 {ev.get('k')} 份）")
    return str(ev.get("reason") or "")


def fidelity_lines(asr, dropped_segments: Optional[list] = None) -> list[str]:
    """Every fidelity annotation as one readable line.

    Shared by the exported .md/.txt tail and the GUI tab so the file and the
    screen can never drift apart. Qt-free on purpose.
    """
    out: list[str] = []
    for a in asr.annotations:
        kind = a.get("type")
        ts = f"{_ts_short(a.get('start', 0))}–{_ts_short(a.get('end', 0))}"
        if kind == "repeat_arbitration":
            verdict = _VERDICT_ZH.get(a.get("verdict"), str(a.get("verdict", "?")))
            state = "已在正文中折叠" if a.get("folded") else "正文原样保留"
            ev = _evidence_brief(a.get("evidence") or {})
            line = f"[{ts}] {verdict} · {state} ｜原话「{a.get('original', '')}」"
            out.append(f"{line}｜依据：{ev}" if ev else line)
        elif kind == "adjacent_duplicate_segment":
            out.append(f"[{ts}] 相邻重复段 · 正文原样保留 ｜"
                       f"「{a.get('text', '')}」与上一段重复")
    for d in (dropped_segments or []):
        ts = f"{_ts_short(d.get('start', 0))}–{_ts_short(d.get('end', 0))}"
        out.append(f"[{ts}] 相邻重复段 · 已整段剔除 ｜"
                   f"「{d.get('text', '')}」与上一段重复（原话见 meta.json）")
    return out


def export_all(
    asr: ASRResult,
    out_dir: Path,
    base: str,
    formats: list[str],
    labels: Optional[dict[int, str]] = None,
    extra_markdown: Optional[str] = None,
    extra_markdown_suffix: str = "report",
    dropped_segments: Optional[list] = None,
) -> dict[str, Path]:
    out_dir = Path(out_dir)
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
        # probe writability (macOS TCC may block protected folders)
        probe = out_dir / ".write_test"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink()
    except (PermissionError, OSError):
        from core.paths import data_root
        fallback = data_root() / base
        fallback.mkdir(parents=True, exist_ok=True)
        out_dir = fallback

    outputs: dict[str, Path] = {}

    for fmt in formats:
        try:
            if fmt == "txt":
                outputs["txt"] = _txt(asr, out_dir, base, labels, dropped_segments)
            elif fmt == "md":
                outputs["md"] = _md(asr, out_dir, base, labels, dropped_segments)
            elif fmt == "docx":
                outputs["docx"] = _docx(asr, out_dir, base, labels, extra_markdown)
            elif fmt == "doc":
                outputs["doc"] = _doc(asr, out_dir, base, labels, extra_markdown)
            elif fmt == "srt":
                outputs["srt"] = _srt(asr, out_dir, base, labels)
            elif fmt == "json":
                outputs["json"] = _json(asr, out_dir, base, labels)
        except Exception as exc:  # one bad writer shouldn't sink the rest
            import logging
            logging.getLogger(__name__).warning("export %s failed: %s", fmt, exc)

    if extra_markdown:
        p = out_dir / f"{base}.{extra_markdown_suffix}.md"
        p.write_text(extra_markdown, encoding="utf-8")
        outputs["report"] = p

    return outputs


def _label(labels, seg_id) -> str:
    if not labels:
        return ""
    raw = labels.get(seg_id)
    if not raw:
        return ""
    return _SPEAKER_ZH.get(raw, raw)


def _txt(asr, out_dir, base, labels, dropped_segments=None) -> Path:
    lines = []
    for s in asr.segments:
        ts = _ts_short(s.start)
        spk = _label(labels, s.id)
        prefix = f"[{ts}] " + (f"{spk}: " if spk else "")
        lines.append(prefix + s.text)
    body = "\n".join(lines) if lines else asr.full_text
    fidelity = fidelity_lines(asr, dropped_segments)
    if fidelity:
        # Appended after the transcript, never interleaved into it.
        body += "\n\n" + "\n".join(
            [FIDELITY_HEADING, "", FIDELITY_NOTE, ""] + fidelity)
    p = out_dir / f"{base}.txt"
    p.write_text(body, encoding="utf-8")
    return p


def _md(asr, out_dir, base, labels, dropped_segments=None) -> Path:
    L = [f"# {base}", "", f"- 语言：{asr.language}", f"- 引擎：{asr.model_used}",
         f"- 时长：{asr.audio_duration_sec:.0f}s", "", "## 转写", ""]
    for s in asr.segments:
        ts = _ts_short(s.start)
        spk = _label(labels, s.id)
        if spk:
            L.append(f"**{spk}** `[{ts}]` {s.text}")
        else:
            L.append(f"`[{ts}]` {s.text}")
        L.append("")
    fidelity = fidelity_lines(asr, dropped_segments)
    if fidelity:
        L.extend(["", f"## {FIDELITY_HEADING}", "", f"> {FIDELITY_NOTE}", ""])
        L.extend(f"- {line}" for line in fidelity)
        L.append("")
    p = out_dir / f"{base}.md"
    p.write_text("\n".join(L), encoding="utf-8")
    return p


def _srt(asr, out_dir, base, labels) -> Path:
    blocks = []
    for i, s in enumerate(asr.segments, 1):
        spk = _label(labels, s.id)
        text = (f"{spk}: " if spk else "") + s.text
        blocks.append(f"{i}\n{_ts_srt(s.start)} --> {_ts_srt(s.end)}\n{text}\n")
    p = out_dir / f"{base}.srt"
    p.write_text("\n".join(blocks), encoding="utf-8")
    return p


def _json(asr, out_dir, base, labels) -> Path:
    data = {
        "language": asr.language,
        "model": asr.model_used,
        "duration_sec": asr.audio_duration_sec,
        "full_text": asr.full_text,
        "segments": [
            {
                "id": s.id,
                "start": round(s.start, 3),
                "end": round(s.end, 3),
                "speaker": labels.get(s.id) if labels else None,
                "text": s.text,
                "avg_logprob": round(s.confidence, 3),
                "words": [
                    {"word": w.word, "start": round(w.start, 3),
                     "end": round(w.end, 3), "confidence": round(w.confidence, 3)}
                    for w in s.words
                ],
            }
            for s in asr.segments
        ],
        "warnings": asr.warnings,
        "annotations": asr.annotations,
    }
    p = out_dir / f"{base}.json"
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


# ----------------------------------------------------------------------
# Word documents
# ----------------------------------------------------------------------

def _build_docx_document(asr, base, labels, report_md):
    """Return a python-docx Document. For IELTS (report_md given) the doc IS the
    report (which already contains the verbatim transcript); otherwise it's the
    timestamped transcript."""
    from docx import Document

    doc = Document()
    if report_md:
        _md_into_docx(doc, report_md)
    else:
        doc.add_heading(base, level=0)
        meta = doc.add_paragraph()
        meta.add_run(f"语言 {asr.language} · 引擎 {asr.model_used} · "
                     f"时长 {asr.audio_duration_sec:.0f}s").italic = True
        for s in asr.segments:
            ts = _ts_short(s.start)
            spk = _label(labels, s.id)
            p = doc.add_paragraph()
            p.add_run(f"[{ts}] ").bold = True
            if spk:
                p.add_run(f"{spk}: ").bold = True
            p.add_run(s.text)
    return doc


def _md_into_docx(doc, md: str) -> None:
    """Render the lightweight markdown our reports use into a docx Document."""
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.style = doc.styles["Intense Quote"] if "Intense Quote" in [s.name for s in doc.styles] else p.style
        elif line.lstrip().startswith("- "):
            doc.add_paragraph(line.lstrip()[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)


def _docx(asr, out_dir, base, labels, report_md) -> Path:
    doc = _build_docx_document(asr, base, labels, report_md)
    p = out_dir / f"{base}.docx"
    doc.save(str(p))
    return p


def _doc(asr, out_dir, base, labels, report_md) -> Path:
    """Legacy .doc via macOS `textutil`, converting from a generated .docx."""
    import shutil
    import subprocess
    import tempfile

    # Build a docx first (reuse formatting), then convert.
    doc = _build_docx_document(asr, base, labels, report_md)
    tmp_docx = Path(tempfile.mkstemp(suffix=".docx", prefix="recorder_")[1])
    doc.save(str(tmp_docx))

    out = out_dir / f"{base}.doc"
    textutil = shutil.which("textutil")
    if not textutil:
        raise RuntimeError(t("export_textutil_missing"))
    proc = subprocess.run(
        [textutil, "-convert", "doc", str(tmp_docx), "-output", str(out)],
        capture_output=True, text=True, check=False,
    )
    tmp_docx.unlink(missing_ok=True)
    if proc.returncode != 0 or not out.exists():
        raise RuntimeError(
            t("export_textutil_failed", err=proc.stderr.strip()[:160]))
    return out


def _ts_short(sec: float) -> str:
    td = timedelta(seconds=int(sec))
    m, s = divmod(int(td.total_seconds()), 60)
    h, m = divmod(m, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:d}:{s:02d}"


def _ts_srt(sec: float) -> str:
    ms = int((sec - int(sec)) * 1000)
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
